#!/usr/bin/env python3
"""
Build a .pptx presentation from a JSON spec.

Spec format:
{
  "title": "Presentation Title",
  "theme": {
    "font_name": "Arial",
    "font_size": 18,
    "background_color": "FFFFFF",
    "heading_color": "000000",
    "body_color": "000000"
  },
  "slides": [
    {
      "title": "Slide Title",
      "subtitle": "Optional subtitle",
      "content": [
        {"type": "bullets", "items": ["item 1", "item 2"]},
        {"type": "table", "header": ["A", "B"], "rows": [["1", "2"]]},
        {"type": "para", "text": "paragraph text"}
      ]
    }
  ]
}
"""

import json
import sys
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, font_name, font_size, bold, color_hex):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = hex_to_rgb(color_hex)


def set_paragraph_font(para, font_name, font_size, bold, color_hex, alignment=None):
    for run in para.runs:
        set_run_font(run, font_name, font_size, bold, color_hex)
    if alignment:
        para.alignment = alignment


def add_slide_textbox(slide, left, top, width, height, text, font_name, font_size, bold, color_hex, alignment=None):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, font_name, font_size, bold, color_hex)
    if alignment:
        p.alignment = alignment
    return txBox


def build(spec, out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    theme = spec.get("theme", {})
    font_name = theme.get("font_name", "Arial")
    font_size = theme.get("font_size", 18)
    bg_color = theme.get("background_color", "FFFFFF")
    heading_color = theme.get("heading_color", "000000")
    body_color = theme.get("body_color", "000000")

    # Use blank layout
    blank_layout = prs.slide_layouts[6]

    for slide_spec in spec.get("slides", []):
        slide = prs.slides.add_slide(blank_layout)

        # Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(bg_color)

        y = 0.5

        # Title
        if "title" in slide_spec:
            add_slide_textbox(slide, 0.7, y, 12, 1.0, slide_spec["title"],
                             font_name, 32, True, heading_color, PP_ALIGN.LEFT)
            y += 1.1

        # Subtitle
        if "subtitle" in slide_spec:
            add_slide_textbox(slide, 0.7, y, 12, 0.6, slide_spec["subtitle"],
                             font_name, 20, False, body_color, PP_ALIGN.LEFT)
            y += 0.7

        # Content
        for block in slide_spec.get("content", []):
            btype = block.get("type")
            if btype == "bullets":
                items = block.get("items", [])
                for item in items:
                    add_slide_textbox(slide, 1.0, y, 11, 0.5, f"• {item}",
                                     font_name, font_size, False, body_color, PP_ALIGN.LEFT)
                    y += 0.45
                y += 0.2
            elif btype == "table":
                header = block.get("header", [])
                rows = block.get("rows", [])
                if header:
                    cols = len(header)
                    table_shape = slide.shapes.add_table(
                        len(rows) + 1, cols,
                        Inches(0.7), Inches(y),
                        Inches(11.5), Inches(min(0.4 * (len(rows) + 1), 4.0))
                    )
                    table = table_shape.table
                    for j, h in enumerate(header):
                        cell = table.cell(0, j)
                        cell.text = h
                        for p in cell.text_frame.paragraphs:
                            set_paragraph_font(p, font_name, font_size - 2, True, heading_color, PP_ALIGN.CENTER)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = hex_to_rgb("F0F0F0")
                    for i, row in enumerate(rows):
                        for j, cell_text in enumerate(row):
                            cell = table.cell(i + 1, j)
                            cell.text = str(cell_text)
                            for p in cell.text_frame.paragraphs:
                                set_paragraph_font(p, font_name, font_size - 2, False, body_color, PP_ALIGN.LEFT)
                            if i % 2 == 1:
                                cell.fill.solid()
                                cell.fill.fore_color.rgb = hex_to_rgb("F5F5F5")
                    y += 0.4 * (len(rows) + 1) + 0.3
            elif btype == "para":
                text = block.get("text", "")
                add_slide_textbox(slide, 0.7, y, 11.5, 1.0, text,
                                 font_name, font_size, False, body_color, PP_ALIGN.LEFT)
                y += 0.6

    prs.save(out_path)
    return out_path


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 1
    with open(sys.argv[1], encoding="utf-8") as fh:
        spec = json.load(fh)
    out = build(spec, sys.argv[2])
    print(json.dumps({"written": out, "slides": len(spec.get("slides", []))}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())